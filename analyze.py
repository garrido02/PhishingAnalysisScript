import os
import openpyxl
import pandas as pd


def extract_campaign_info(filename):
    """Extracts the location and the type of campaign from the filename.

    Example: 2025C2_PTAOMZ_D_LCGemini2.csv -> Loc: PTAOMZ, Type: LC
    LC = Credentials, A = Anex, LCA = Credentials & Anex
    """
    clean_name = os.path.splitext(filename)[0]
    parts = clean_name.split("_")

    location = "Unknown"
    camp_type = "Unknown"

    if len(parts) >= 4:
        location = parts[1]
        type_str = parts[3].upper()

        # Checking 'LCA' first is mandatory, otherwise 'LC' would intercept it
        if type_str.startswith("LCA"):
            camp_type = "Credentials & Anex"
        elif type_str.startswith("LC"):
            camp_type = "Credentials"
        elif type_str.startswith("A"):
            camp_type = "Anex"

    return location, camp_type


def generate_report(master_file_path):
    """Reads the output file and calculates both base metrics and percentages,

    saving the results in separate tables in 'Analytics' and 'Analysis - Users' sheets.
    """
    print("\n[*] Starting analytics generation...")

    if not os.path.exists(master_file_path):
        print(
            f"   [ERROR] Master file {master_file_path} not found. Cannot generate report."
        )
        return

    try:
        df = pd.read_excel(master_file_path, sheet_name=0)

        if "Source File" not in df.columns:
            print(
                "   [ERROR] Column 'Source File' not found. Cannot extract campaign info."
            )
            return

        # Find UTC variation in headers to determine which columns to use for action flags
        utc_flag = any("UTC" in header for header in df.columns)

        # Identify the user identification column dynamically
        user_col = None
        for col in ["Email", "Recipient", "User", "Utilizador"]:
            if col in df.columns:
                user_col = col
                break

        if not user_col:
            print(
                "   [ERROR] User column (Email/Recipient) not found. Cannot generate user analytics."
            )
            return

        # Extract the location and campaign type from the 'Source File' column
        df[["Location_Extracted", "Campaign_Type"]] = df["Source File"].apply(
            lambda x: pd.Series(extract_campaign_info(x))
        )

        # Base action flags from source columns
        if utc_flag:
            df["is_opened"] = (
                df["Opened At (UTC)"].notna().astype(int)
                if "Opened At (UTC)" in df.columns
                else 0
            )
            df["is_clicked"] = (
                df["Clicked At (UTC)"].notna().astype(int)
                if "Clicked At (UTC)" in df.columns
                else 0
            )
            df["is_data_entered"] = (
                df["Data Entered At (UTC)"].notna().astype(int)
                if "Data Entered At (UTC)" in df.columns
                else 0
            )
            df["is_attachment_opened"] = (
                df["Attachment Opened At (UTC)"].notna().astype(int)
                if "Attachment Opened At (UTC)" in df.columns
                else 0
            )
        else:
            df["is_opened"] = (
                df["Opened At"].notna().astype(int) if "Opened At" in df.columns else 0
            )
            df["is_clicked"] = (
                df["Clicked At"].notna().astype(int)
                if "Clicked At" in df.columns
                else 0
            )
            df["is_data_entered"] = (
                df["Data Entered At"].notna().astype(int)
                if "Data Entered At" in df.columns
                else 0
            )
            df["is_attachment_opened"] = (
                df["Attachment Opened At"].notna().astype(int)
                if "Attachment Opened At" in df.columns
                else 0
            )

        # Strict separation flags for campaign typologies
        df["type_pure_LC"] = (df["Campaign_Type"] == "Credentials").astype(int)
        df["type_pure_A"] = (df["Campaign_Type"] == "Anex").astype(int)
        df["type_LCA"] = (df["Campaign_Type"] == "Credentials & Anex").astype(int)

        # Links Clicked breakdown (LCA data entry implicitly means they clicked the link)
        df["click_LC"] = ((df["type_pure_LC"] == 1) & (df["is_clicked"] == 1)).astype(
            int
        )
        df["click_A"] = ((df["type_pure_A"] == 1) & (df["is_clicked"] == 1)).astype(int)
        df["click_LCA"] = (
            (df["type_LCA"] == 1)
            & ((df["is_clicked"] == 1) | (df["is_data_entered"] == 1))
        ).astype(int)
        df["click_total"] = (
            (df["click_LC"] == 1) | (df["click_A"] == 1) | (df["click_LCA"] == 1)
        ).astype(int)

        # Credentials Entered breakdown
        df["insert_LC"] = (
            (df["type_pure_LC"] == 1) & (df["is_data_entered"] == 1)
        ).astype(int)
        df["insert_A"] = (
            (df["type_pure_A"] == 1) & (df["is_data_entered"] == 1)
        ).astype(int)
        df["insert_LCA"] = (
            (df["type_LCA"] == 1) & (df["is_data_entered"] == 1)
        ).astype(int)
        df["insert_total"] = (
            (df["insert_LC"] == 1) | (df["insert_A"] == 1) | (df["insert_LCA"] == 1)
        ).astype(int)

        # Annexes Opened breakdown
        df["attachment_A"] = (
            (df["type_pure_A"] == 1) & (df["is_attachment_opened"] == 1)
        ).astype(int)
        df["attachment_LCA"] = (
            (df["type_LCA"] == 1) & (df["is_attachment_opened"] == 1)
        ).astype(int)
        df["attachment_total"] = (
            (df["attachment_A"] == 1) | (df["attachment_LCA"] == 1)
        ).astype(int)

        # TABELA 1: LOCATION BASE METRICS TABLE (Purely Campaign Volumes)
        metrics = df.groupby("Location_Extracted").apply(
            lambda g: pd.Series(
                {
                    "Total Sent Emails": len(g),
                    "Email w/Credentials (LC)": g["type_pure_LC"].sum(),
                    "Email w/Anex (A)": g["type_pure_A"].sum(),
                    "Email w/Both (LCA)": g["type_LCA"].sum(),
                    "Email Opened (Total)": g["is_opened"].sum(),
                    "Emails Opened (LC)": (
                        (g["type_pure_LC"] == 1) & (g["is_opened"] == 1)
                    ).sum(),
                    "Emails Opened (A)": (
                        (g["type_pure_A"] == 1) & (g["is_opened"] == 1)
                    ).sum(),
                    "Emails Opened (LCA)": (
                        (g["type_LCA"] == 1) & (g["is_opened"] == 1)
                    ).sum(),
                    "Links Clicked (Total)": g["click_total"].sum(),
                    "Credentials Entered (Total)": g["insert_total"].sum(),
                    "Annexes Opened (Total)": g["attachment_total"].sum(),
                    "Links Clicked (LC)": g["click_LC"].sum(),
                    "Links Clicked (A)": g["click_A"].sum(),
                    "Links Clicked (LCA)": g["click_LCA"].sum(),
                    "Credentials Entered (LC)": g["insert_LC"].sum(),
                    "Credentials Entered (A)": g["insert_A"].sum(),
                    "Credentials Entered (LCA)": g["insert_LCA"].sum(),
                    "Annexes Opened (A)": g["attachment_A"].sum(),
                    "Annexes Opened (LCA)": g["attachment_LCA"].sum(),
                }
            )
        )

        total_row = metrics.sum().rename("TOTAL GLOBAL")
        report_df = pd.concat([metrics, total_row.to_frame().T])

        # TABELA 2: LOCATION PERCENTAGES TABLE (Purely Campaign Conversion Rates %)
        pct_df = pd.DataFrame(index=report_df.index)

        def calc_pct(num_col, den_col):
            calc = report_df[num_col] / report_df[den_col].replace(0, pd.NA)
            return calc.fillna(0).astype(float)

        # 1. Global Performance
        pct_df["% Opened (Global)"] = calc_pct(
            "Email Opened (Total)", "Total Sent Emails"
        )
        pct_df["% Clicked (from Opened Global)"] = calc_pct(
            "Links Clicked (Total)", "Email Opened (Total)"
        )
        pct_df["% Credentials Entered (from Clicked Global)"] = calc_pct(
            "Credentials Entered (Total)", "Links Clicked (Total)"
        )
        pct_df["% Annexes Opened (from Clicked Global)"] = calc_pct(
            "Annexes Opened (Total)", "Links Clicked (Total)"
        )

        # 2. Pure LC Performance
        pct_df["% Opened (LC)"] = calc_pct(
            "Emails Opened (LC)", "Email w/Credentials (LC)"
        )
        pct_df["% Clicked (from Opened LC)"] = calc_pct(
            "Links Clicked (LC)", "Emails Opened (LC)"
        )
        pct_df["% Credentials Entered (from Clicked LC)"] = calc_pct(
            "Credentials Entered (LC)", "Links Clicked (LC)"
        )

        # 3. Pure A Performance
        pct_df["% Opened (A)"] = calc_pct("Emails Opened (A)", "Email w/Anex (A)")
        pct_df["% Clicked (from Opened A)"] = calc_pct(
            "Links Clicked (A)", "Emails Opened (A)"
        )
        pct_df["% Annexes Opened (from Clicked A)"] = calc_pct(
            "Annexes Opened (A)", "Links Clicked (A)"
        )

        # 4. LCA Performance
        pct_df["% Opened (LCA)"] = calc_pct("Emails Opened (LCA)", "Email w/Both (LCA)")
        pct_df["% Clicked (from Opened LCA)"] = calc_pct(
            "Links Clicked (LCA)", "Emails Opened (LCA)"
        )
        pct_df["% Credentials Entered (from Clicked LCA)"] = calc_pct(
            "Credentials Entered (LCA)", "Links Clicked (LCA)"
        )
        pct_df["% Annexes Opened (from Clicked LCA)"] = calc_pct(
            "Annexes Opened (LCA)", "Links Clicked (LCA)"
        )

        # Map unique user behavior per location to guarantee mathematical alignment
        user_loc_summary = (
            df.groupby(["Location_Extracted", user_col])
            .agg(
                has_opened=("is_opened", "max"),
                has_clicked=("click_total", "max"),
                has_inserted=("insert_total", "max"),
                has_attached=("attachment_total", "max"),
                total_clicks=("click_total", "sum"),
            )
            .reset_index()
        )

        # TABELA 3: USER METRICS & CLICK DISTRIBUTION VOLUMES
        dist_by_loc = user_loc_summary.groupby("Location_Extracted").apply(
            lambda g: pd.Series(
                {
                    "Total Users": len(g),
                    "Users Opened (Total)": g["has_opened"].sum(),
                    "Users Clicked (Total)": g["has_clicked"].sum(),
                    "Users Entered Credentials (Total)": g["has_inserted"].sum(),
                    "Users Opened Annex (Total)": g["has_attached"].sum(),
                    "Users with 0 clicks": (g["total_clicks"] == 0).sum(),
                    "Users with 1 click": (g["total_clicks"] == 1).sum(),
                    "Users with 2 clicks": (g["total_clicks"] == 2).sum(),
                    "Users with 3+ clicks": (g["total_clicks"] >= 3).sum(),
                }
            )
        )

        dist_total = dist_by_loc.sum().rename("TOTAL GLOBAL")
        dist_by_loc_report = pd.concat([dist_by_loc, dist_total.to_frame().T])

        # TABELA 4: USER METRICS & CLICK DISTRIBUTION RATES %
        dist_pct_df = pd.DataFrame(index=dist_by_loc_report.index)
        den_dist = dist_by_loc_report["Total Users"].replace(0, pd.NA)
        den_opened_users = dist_by_loc_report["Users Opened (Total)"].replace(0, pd.NA)
        den_clicked_users = dist_by_loc_report["Users Clicked (Total)"].replace(
            0, pd.NA
        )

        # Unique User Funnel Rates
        dist_pct_df["% Users Opened (from Total Users)"] = (
            (dist_by_loc_report["Users Opened (Total)"] / den_dist)
            .fillna(0)
            .astype(float)
        )
        dist_pct_df["% Users Clicked (from Opened Users)"] = (
            (dist_by_loc_report["Users Clicked (Total)"] / den_opened_users)
            .fillna(0)
            .astype(float)
        )
        dist_pct_df["% Users Entered Credentials (from Clicked Users)"] = (
            (
                dist_by_loc_report["Users Entered Credentials (Total)"]
                / den_clicked_users
            )
            .fillna(0)
            .astype(float)
        )
        dist_pct_df["% Users Opened Annex (from Clicked Users)"] = (
            (dist_by_loc_report["Users Opened Annex (Total)"] / den_clicked_users)
            .fillna(0)
            .astype(float)
        )

        # Click Cohort Percentages relative to the total population of the location
        dist_pct_df["% Users with 0 clicks"] = (
            (dist_by_loc_report["Users with 0 clicks"] / den_dist)
            .fillna(0)
            .astype(float)
        )
        dist_pct_df["% Users with 1 click"] = (
            (dist_by_loc_report["Users with 1 click"] / den_dist)
            .fillna(0)
            .astype(float)
        )
        dist_pct_df["% Users with 2 clicks"] = (
            (dist_by_loc_report["Users with 2 clicks"] / den_dist)
            .fillna(0)
            .astype(float)
        )
        dist_pct_df["% Users with 3+ clicks"] = (
            (dist_by_loc_report["Users with 3+ clicks"] / den_dist)
            .fillna(0)
            .astype(float)
        )

        user_metrics = df.groupby(user_col).apply(
            lambda g: pd.Series(
                {
                    "Total Sent Emails": len(g),
                    "Email w/Credentials (LC)": g["type_pure_LC"].sum(),
                    "Email w/Anex (A)": g["type_pure_A"].sum(),
                    "Email w/Both (LCA)": g["type_LCA"].sum(),
                    "Email Opened (Total)": g["is_opened"].sum(),
                    "Emails Opened (LC)": (
                        (g["type_pure_LC"] == 1) & (g["is_opened"] == 1)
                    ).sum(),
                    "Emails Opened (A)": (
                        (g["type_pure_A"] == 1) & (g["is_opened"] == 1)
                    ).sum(),
                    "Emails Opened (LCA)": (
                        (g["type_LCA"] == 1) & (g["is_opened"] == 1)
                    ).sum(),
                    "Links Clicked (Total)": g["click_total"].sum(),
                    "Credentials Entered (Total)": g["insert_total"].sum(),
                    "Annexes Opened (Total)": g["attachment_total"].sum(),
                    "Links Clicked (LC)": g["click_LC"].sum(),
                    "Links Clicked (A)": g["click_A"].sum(),
                    "Links Clicked (LCA)": g["click_LCA"].sum(),
                    "Credentials Entered (LC)": g["insert_LC"].sum(),
                    "Credentials Entered (A)": g["insert_A"].sum(),
                    "Credentials Entered (LCA)": g["insert_LCA"].sum(),
                    "Annexes Opened (A)": g["attachment_A"].sum(),
                    "Annexes Opened (LCA)": g["attachment_LCA"].sum(),
                }
            )
        )

        user_pct_df = pd.DataFrame(index=user_metrics.index)

        def calc_user_pct(num_col, den_col):
            calc = user_metrics[num_col] / user_metrics[den_col].replace(0, pd.NA)
            return calc.fillna(0).astype(float)

        user_pct_df["% Opened (Global)"] = calc_user_pct(
            "Email Opened (Total)", "Total Sent Emails"
        )
        user_pct_df["% Clicked (from Opened Global)"] = calc_user_pct(
            "Links Clicked (Total)", "Email Opened (Total)"
        )
        user_pct_df["% Credentials Entered (from Clicked Global)"] = calc_user_pct(
            "Credentials Entered (Total)", "Links Clicked (Total)"
        )
        user_pct_df["% Annexes Opened (from Clicked Global)"] = calc_user_pct(
            "Annexes Opened (Total)", "Links Clicked (Total)"
        )
        user_pct_df["% Opened (LC)"] = calc_user_pct(
            "Emails Opened (LC)", "Email w/Credentials (LC)"
        )
        user_pct_df["% Clicked (from Opened LC)"] = calc_user_pct(
            "Links Clicked (LC)", "Emails Opened (LC)"
        )
        user_pct_df["% Credentials Entered (from Clicked LC)"] = calc_user_pct(
            "Credentials Entered (LC)", "Links Clicked (LC)"
        )
        user_pct_df["% Opened (A)"] = calc_user_pct(
            "Emails Opened (A)", "Email w/Anex (A)"
        )
        user_pct_df["% Clicked (from Opened A)"] = calc_user_pct(
            "Links Clicked (A)", "Emails Opened (A)"
        )
        user_pct_df["% Annexes Opened (from Clicked A)"] = calc_user_pct(
            "Annexes Opened (A)", "Links Clicked (A)"
        )
        user_pct_df["% Opened (LCA)"] = calc_user_pct(
            "Emails Opened (LCA)", "Email w/Both (LCA)"
        )
        user_pct_df["% Clicked (from Opened LCA)"] = calc_user_pct(
            "Links Clicked (LCA)", "Emails Opened (LCA)"
        )
        user_pct_df["% Annexes Opened (from Clicked LCA)"] = calc_user_pct(
            "Annexes Opened (LCA)", "Links Clicked (LCA)"
        )
        user_pct_df["% Credentials Entered (from Clicked LCA)"] = calc_user_pct(
            "Credentials Entered (LCA)", "Links Clicked (LCA)"
        )

        # EXCEL EXPORT LOGIC
        book = openpyxl.load_workbook(master_file_path)
        if "Analytics" in book.sheetnames:
            del book["Analytics"]
        if "Analysis - Users" in book.sheetnames:
            del book["Analysis - Users"]
        book.save(master_file_path)
        book.close()

        with pd.ExcelWriter(
            master_file_path, engine="openpyxl", mode="a", if_sheet_exists="overlay"
        ) as writer:
            # Table 1: Write base metrics (Pure Campaign Volumes)
            report_df.to_excel(
                writer, sheet_name="Analytics", index_label="Location (Volumes)"
            )

            # Table 2: Write conversion percentages table below metrics
            startrow_pct = len(report_df) + 4
            pct_df.to_excel(
                writer,
                sheet_name="Analytics",
                startrow=startrow_pct,
                index_label="Location (Conversion Rates %)",
            )

            # Table 3: Write user click distribution matrix volumes grouped by location
            startrow_dist = startrow_pct + len(pct_df) + 4
            dist_by_loc_report.to_excel(
                writer,
                sheet_name="Analytics",
                startrow=startrow_dist,
                index_label="Location (User Click Distribution Volumes)",
            )

            # Table 4: Write user click distribution percentages table below distribution volumes
            startrow_dist_pct = startrow_dist + len(dist_by_loc_report) + 4
            dist_pct_df.to_excel(
                writer,
                sheet_name="Analytics",
                startrow=startrow_dist_pct,
                index_label="Location (User Click Distribution Rates %)",
            )

            # Format the cell masks natively for the percentage tables on Sheet 1
            worksheet = writer.sheets["Analytics"]

            # Mask Table 2: Location Conversion Rates Table
            excel_start_row = startrow_pct + 2
            excel_end_row = startrow_pct + 1 + len(pct_df)
            excel_start_col = 2
            excel_end_col = 1 + len(pct_df.columns)

            for row in range(excel_start_row, excel_end_row + 1):
                for col in range(excel_start_col, excel_end_col + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.number_format = "0.00%"

            # Mask Table 4: User Click Distribution Rates Table
            excel_dpct_start_row = startrow_dist_pct + 2
            excel_dpct_end_row = startrow_dist_pct + 1 + len(dist_pct_df)
            excel_dpct_start_col = 2
            excel_dpct_end_col = 1 + len(dist_pct_df.columns)

            for row in range(excel_dpct_start_row, excel_dpct_end_row + 1):
                for col in range(excel_dpct_start_col, excel_dpct_end_col + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.number_format = "0.00%"

            # A. Write individual user base metrics
            user_metrics.to_excel(
                writer, sheet_name="Analysis - Users", index_label="User (Volumes)"
            )

            # B. Write individual conversion percentages below user metrics
            startrow_user_pct = len(user_metrics) + 4
            user_pct_df.to_excel(
                writer,
                sheet_name="Analysis - Users",
                startrow=startrow_user_pct,
                index_label="User (Conversion Rates %)",
            )

            # Format user percentage columns natively
            user_worksheet = writer.sheets["Analysis - Users"]
            excel_user_start_row = startrow_user_pct + 2
            excel_user_end_row = startrow_user_pct + 1 + len(user_pct_df)
            excel_user_start_col = 2
            excel_user_end_col = 1 + len(user_pct_df.columns)

            for row in range(excel_user_start_row, excel_user_end_row + 1):
                for col in range(excel_user_start_col, excel_user_end_col + 1):
                    cell = user_worksheet.cell(row=row, column=col)
                    cell.number_format = "0.00%"

        print(f"   [+] Data written to sheet 'Analytics' in {master_file_path}")
        print(
            f"   [+] User data written to sheet 'Analysis - Users' in {master_file_path}"
        )

    except Exception as e:
        print(f"   [ERROR] Failed to generate analytics: {e}")
